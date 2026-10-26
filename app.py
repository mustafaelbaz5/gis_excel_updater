from __future__ import annotations

from collections import Counter, defaultdict
from copy import deepcopy
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any
import math
import re
import unicodedata
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

BASE_TO_GIS = {
    "اسم الحوض": "Basin_Name",
    "كود الحوض": "Basin_Code",
    "فدان": "fadan_survey",
    "قراط": "qirat_survey",
    "سهم": "sahm_survey",
    "نوع الاستخدام": "using_type",
    "نوع المحصول": "Crop_type",
    "مراحل النمو": "stage_of_growth",
    "ملاحظات على الحيازة من فريق العمل": "notes_from_survey",
}
KEY_NAME = "كود القطعة"
SKIP = "— لا يتم التحديث —"
BG, CARD, INK, MUTED, PRIMARY, OK, BAD, WARN = (
    "#F4F7FA", "#FFFFFF", "#16233A", "#617083", "#176B87", "#14804A", "#B42318", "#B54708"
)


class DataError(Exception):
    pass


@dataclass
class Table:
    path: Path
    headers: list[str]
    rows: list[dict[str, Any]]
    workbook: Any = None
    sheet: Any = None
    native: bool = False
    header_row: int = 1


@dataclass
class Report:
    base_total: int = 0
    gis_total: int = 0
    matched: int = 0
    updated_rows: int = 0
    changes: Counter = field(default_factory=Counter)
    blank_changes: Counter = field(default_factory=Counter)
    rule_changes: int = 0
    base_unmatched: list[str] = field(default_factory=list)
    gis_unmatched: list[str] = field(default_factory=list)
    base_duplicates: list[str] = field(default_factory=list)
    gis_duplicates: list[str] = field(default_factory=list)
    base_blank_keys: int = 0
    gis_blank_keys: int = 0
    details: list[tuple[str, str, Any, Any, str]] = field(default_factory=list)

    @property
    def changed_cells(self) -> int:
        return sum(self.changes.values())


def text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).replace("\u200f", "").replace("\u200e", "").strip()


def norm(value: Any) -> str:
    value = unicodedata.normalize("NFKC", text(value))
    value = value.translate(str.maketrans({"أ": "ا", "إ": "ا", "آ": "ا", "ى": "ي"}))
    return re.sub(r"\s+", " ", value).casefold()


def blank(value: Any) -> bool:
    return value is None or (isinstance(value, float) and math.isnan(value)) or text(value).casefold() in {"", "null"}


def key(value: Any) -> str | None:
    if blank(value):
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    value = text(value).translate(str.maketrans("٠١٢٣٤٥٦٧٨٩", "0123456789"))
    if re.fullmatch(r"[+-]?\d+\.0+", value):
        value = value.split(".", 1)[0]
    return value or None


def equal(left: Any, right: Any) -> bool:
    if blank(left) and blank(right):
        return True
    if isinstance(left, (int, float)) and isinstance(right, (int, float)):
        return float(left) == float(right)
    return text(left) == text(right)


def first_header(headers: list[str], wanted: str) -> str | None:
    target = norm(wanted)
    return next((item for item in headers if norm(item) == target), None)


def header_row(matrix: list[list[Any]]) -> int:
    if not matrix:
        raise DataError("الملف لا يحتوي على بيانات.")
    counts = [sum(not blank(value) for value in row) for row in matrix[:20]]
    best = max(range(len(counts)), key=counts.__getitem__)
    if counts[best] == 0:
        raise DataError("لم يتم العثور على صف عناوين للأعمدة.")
    return best


def unique_headers(raw: list[Any]) -> list[str]:
    result, seen = [], set()
    for index, value in enumerate(raw, 1):
        base = text(value) or f"عمود_{index}"
        candidate, number = base, 2
        while norm(candidate) in seen:
            candidate = f"{base} ({number})"
            number += 1
        result.append(candidate)
        seen.add(norm(candidate))
    return result


def one_openpyxl_sheet(workbook):
    candidates = []
    for sheet in workbook.worksheets:
        if sheet.sheet_state != "visible":
            continue
        sample = [list(row) for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 20), values_only=True)]
        if any(sum(not blank(cell) for cell in row) >= 2 for row in sample):
            candidates.append(sheet)
    if not candidates:
        raise DataError("لم يتم العثور على شيت بيانات صالح.")
    if len(candidates) != 1:
        raise DataError("الملف يحتوي على أكثر من شيت به بيانات. الأداة تعمل على شيت بيانات واحد فقط.")
    return candidates[0]


def read_native(path: Path) -> Table:
    from openpyxl import load_workbook

    workbook = load_workbook(path, keep_vba=path.suffix.lower() == ".xlsm", data_only=False)
    sheet = one_openpyxl_sheet(workbook)
    matrix = [list(row) for row in sheet.iter_rows(values_only=True)]
    start = header_row(matrix)
    headers = unique_headers(matrix[start])
    rows = []
    for excel_row, values in enumerate(matrix[start + 1:], start=start + 2):
        if not any(not blank(value) for value in values):
            continue
        row = {name: values[index] if index < len(values) else None for index, name in enumerate(headers)}
        row["__excel_row__"] = excel_row
        rows.append(row)
    return Table(path, headers, rows, workbook, sheet, True, start + 1)


def read_tabular(path: Path) -> Table:
    import pandas as pd

    try:
        if path.suffix.lower() in {".csv", ".tsv"}:
            frame = pd.read_csv(path, header=None, sep="\t" if path.suffix.lower() == ".tsv" else None, engine="python")
        else:
            excel = pd.ExcelFile(path)
            if len(excel.sheet_names) != 1:
                raise DataError("الملف يحتوي على أكثر من شيت. استخدم ملفًا به شيت بيانات واحد.")
            frame = pd.read_excel(path, header=None)
    except DataError:
        raise
    except Exception as error:
        raise DataError(f"تعذر قراءة الملف: {error}") from error

    matrix = frame.where(frame.notna(), None).values.tolist()
    start = header_row(matrix)
    headers = unique_headers(matrix[start])
    rows = []
    for excel_row, values in enumerate(matrix[start + 1:], start=start + 2):
        if not any(not blank(value) for value in values):
            continue
        row = {name: values[index] if index < len(values) else None for index, name in enumerate(headers)}
        row["__excel_row__"] = excel_row
        rows.append(row)
    return Table(path, headers, rows)


def read_file(value: str) -> Table:
    path = Path(value)
    if not path.exists():
        raise DataError("الملف المحدد غير موجود.")
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        return read_native(path)
    if suffix in {".xls", ".xlsb", ".csv", ".tsv"}:
        return read_tabular(path)
    raise DataError("الصيغة غير مدعومة. اختر xlsx أو xls أو xlsm أو xlsb أو csv أو tsv.")


def defaults(base: Table, gis: Table):
    mapping = {}
    for target, source in BASE_TO_GIS.items():
        mapping[first_header(base.headers, target) or target] = first_header(gis.headers, source)
    return mapping, first_header(base.headers, KEY_NAME), first_header(gis.headers, KEY_NAME)


def index_rows(table: Table, key_header: str):
    result = defaultdict(list)
    empty = 0
    for row in table.rows:
        value = key(row.get(key_header))
        if value is None:
            empty += 1
        else:
            result[value].append(row)
    return result, empty


def write_value(table: Table, row: dict[str, Any], header: str, value: Any):
    row[header] = value
    if table.native:
        table.sheet.cell(row=row["__excel_row__"], column=table.headers.index(header) + 1).value = value


def apply_update(base: Table, gis: Table, base_key: str, gis_key: str, mapping: dict[str, str | None]) -> Report:
    if base_key not in base.headers or gis_key not in gis.headers:
        raise DataError("اختر عمود كود القطعة الصحيح في الملفين.")

    report = Report(len(base.rows), len(gis.rows))
    base_index, report.base_blank_keys = index_rows(base, base_key)
    gis_index, report.gis_blank_keys = index_rows(gis, gis_key)
    base_dupes = {item for item, rows in base_index.items() if len(rows) > 1}
    gis_dupes = {item for item, rows in gis_index.items() if len(rows) > 1}
    report.base_duplicates = sorted(base_dupes)
    report.gis_duplicates = sorted(gis_dupes)
    report.base_unmatched = sorted(item for item in base_index if item not in gis_index and item not in base_dupes)
    report.gis_unmatched = sorted(item for item in gis_index if item not in base_index and item not in gis_dupes)

    active = {target: source for target, source in mapping.items() if target in base.headers and source in gis.headers}
    usage = first_header(base.headers, "نوع الاستخدام")
    crop = first_header(base.headers, "نوع المحصول")
    stage = first_header(base.headers, "مراحل النمو")

    for parcel, base_rows in base_index.items():
        if parcel in base_dupes or parcel in gis_dupes or parcel not in gis_index:
            continue
        base_row, gis_row = base_rows[0], gis_index[parcel][0]
        report.matched += 1
        row_changed = False

        for target, source in active.items():
            old, new = base_row.get(target), (None if blank(gis_row.get(source)) else gis_row.get(source))
            if equal(old, new):
                continue
            write_value(base, base_row, target, new)
            report.changes[target] += 1
            report.blank_changes[target] += int(blank(new))
            report.details.append((parcel, target, old, new, "تحديث من GIS"))
            row_changed = True

        if usage and norm(base_row.get(usage)) in {"بور", "مباني"}:
            for target in (crop, stage):
                if target and not blank(base_row.get(target)):
                    old = base_row.get(target)
                    write_value(base, base_row, target, None)
                    report.changes[target] += 1
                    report.blank_changes[target] += 1
                    report.rule_changes += 1
                    report.details.append((parcel, target, old, None, "قاعدة بور/مباني"))
                    row_changed = True

        if row_changed:
            report.updated_rows += 1
    return report


def output_path(base: Table) -> Path:
    suffix = base.path.suffix.lower() if base.path.suffix.lower() in {".xlsx", ".xlsm"} else ".xlsx"
    candidate = base.path.with_name(f"{base.path.stem}_محدث{suffix}")
    counter = 2
    while candidate.exists():
        candidate = base.path.with_name(f"{base.path.stem}_محدث_{counter}{suffix}")
        counter += 1
    return candidate


def save(base: Table, destination: Path) -> Path:
    if base.native:
        base.workbook.save(destination)
        return destination

    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "البيانات المحدثة"
    sheet.append(base.headers)
    for row in base.rows:
        sheet.append([row.get(header) for header in base.headers])
    sheet.freeze_panes = "A2"
    workbook.save(destination)
    return destination


def report_text(report: Report) -> str:
    lines = [
        "ملخص التحديث",
        f"إجمالي صفوف الملف الأساسي: {report.base_total}",
        f"إجمالي صفوف ملف GIS: {report.gis_total}",
        f"القطع المرتبطة والمؤهلة للتحديث: {report.matched}",
        f"القطع التي تغيّر بها حقل واحد على الأقل: {report.updated_rows}",
        f"إجمالي الخلايا التي تغيّرت: {report.changed_cells}",
        f"تعديلات بسبب قاعدة بور/مباني: {report.rule_changes}",
        "",
        "التغييرات حسب العمود:",
    ]
    if report.changes:
        lines.extend(
            f"- {field}: {count} تحديث، منها {report.blank_changes[field]} تحولت إلى فراغ"
            for field, count in report.changes.items()
        )
    else:
        lines.append("- لا توجد تغييرات")
    lines.extend([
        "",
        f"أكواد الأساسي غير المرتبطة: {len(report.base_unmatched)}",
        f"أكواد GIS غير المرتبطة: {len(report.gis_unmatched)}",
        f"أكواد مكررة في الأساسي: {len(report.base_duplicates)}",
        f"أكواد مكررة في GIS: {len(report.gis_duplicates)}",
        f"صفوف بكود فارغ في الأساسي: {report.base_blank_keys}",
        f"صفوف بكود فارغ في GIS: {report.gis_blank_keys}",
    ])
    return "\n".join(lines)


class ScrollFrame(ttk.Frame):
    def __init__(self, master):
        super().__init__(master, style="Card.TFrame")
        self.canvas = tk.Canvas(self, bg=CARD, highlightthickness=0)
        bar = ttk.Scrollbar(self, orient="vertical", command=self.canvas.yview)
        self.inner = ttk.Frame(self.canvas, style="Card.TFrame")
        self.inner.bind("<Configure>", lambda _event: self.canvas.configure(scrollregion=self.canvas.bbox("all")))
        self.window = self.canvas.create_window((0, 0), window=self.inner, anchor="nw")
        self.canvas.bind("<Configure>", lambda event: self.canvas.itemconfigure(self.window, width=event.width))
        self.canvas.configure(yscrollcommand=bar.set)
        self.canvas.pack(side="left", fill="both", expand=True)
        bar.pack(side="right", fill="y")


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("مُحدِّث بيانات الحيازات من GIS")
        self.geometry("1160x800")
        self.minsize(980, 680)
        self.configure(bg=BG)
        self.option_add("*Font", ("Segoe UI", 10))
        self.base: Table | None = None
        self.gis: Table | None = None
        self.mapping_vars: dict[str, tk.StringVar] = {}
        self.base_key = tk.StringVar()
        self.gis_key = tk.StringVar()
        self.make_style()
        self.build()

    def make_style(self):
        style = ttk.Style(self)
        style.theme_use("clam")
        style.configure("TFrame", background=BG)
        style.configure("Card.TFrame", background=CARD)
        style.configure("Title.TLabel", background=BG, foreground=INK, font=("Segoe UI", 19, "bold"))
        style.configure("Sub.TLabel", background=BG, foreground=MUTED)
        style.configure("Card.TLabel", background=CARD, foreground=INK)
        style.configure("CardTitle.TLabel", background=CARD, foreground=INK, font=("Segoe UI", 11, "bold"))
        style.configure("TButton", padding=(12, 8), font=("Segoe UI", 10, "bold"))
        style.configure("Primary.TButton", background=PRIMARY, foreground="white")
        style.map("Primary.TButton", background=[("active", "#0F526B"), ("disabled", "#A9C5CF")])
        style.configure("Success.TButton", background=OK, foreground="white")
        style.map("Success.TButton", background=[("active", "#0E633B"), ("disabled", "#A3D0B3")])
        style.configure("Treeview", rowheight=28, font=("Segoe UI", 9))
        style.configure("Treeview.Heading", background="#E7F0F4", foreground=INK, font=("Segoe UI", 9, "bold"))
        style.configure("TNotebook", background=BG)
        style.configure("TNotebook.Tab", padding=(15, 7), font=("Segoe UI", 9, "bold"))

    def card(self, parent, pad=16):
        return ttk.Frame(parent, style="Card.TFrame", padding=pad)

    def build(self):
        head = ttk.Frame(self, padding=(24, 20, 24, 6))
        head.pack(fill="x")
        ttk.Label(head, text="مُحدِّث بيانات الحيازات من GIS", style="Title.TLabel", anchor="e").pack(fill="x")
        ttk.Label(head, text="تحديث آمن للملف الأساسي مع تقرير كامل عن الربط والتغييرات.", style="Sub.TLabel", anchor="e").pack(fill="x", pady=(4, 0))

        body = ttk.Frame(self, padding=(24, 8, 24, 20))
        body.pack(fill="both", expand=True)
        files = self.card(body)
        files.pack(fill="x")
        files.columnconfigure(1, weight=1)
        ttk.Label(files, text="1. اختيار الملفين", style="CardTitle.TLabel", anchor="e").grid(row=0, column=0, columnspan=3, sticky="ew", pady=(0, 8))
        self.base_status = ttk.Label(files, text="لم يتم اختيار الملف الأساسي", style="Card.TLabel", anchor="e")
        self.gis_status = ttk.Label(files, text="لم يتم اختيار ملف GIS", style="Card.TLabel", anchor="e")
        for row, label, status, command, button in [
            (1, "الملف الأساسي", self.base_status, self.pick_base, "اختيار الملف الأساسي"),
            (2, "ملف GIS", self.gis_status, self.pick_gis, "اختيار ملف GIS"),
        ]:
            ttk.Label(files, text=label, style="Card.TLabel", anchor="e").grid(row=row, column=0, sticky="e", padx=(0, 12), pady=5)
            status.grid(row=row, column=1, sticky="ew", pady=5)
            ttk.Button(files, text=button, command=command, style="Primary.TButton").grid(row=row, column=2, pady=5)

        self.map_card = self.card(body)
        self.map_card.pack(fill="both", expand=True, pady=(14, 0))
        ttk.Label(self.map_card, text="2. مراجعة ربط الأعمدة", style="CardTitle.TLabel", anchor="e").pack(fill="x")
        self.map_note = ttk.Label(self.map_card, text="اختر الملفين ليظهر الربط التلقائي.", style="Card.TLabel", anchor="e")
        self.map_note.pack(fill="x", pady=(4, 8))
        self.map_area = ScrollFrame(self.map_card)
        self.map_area.pack(fill="both", expand=True)

        controls = ttk.Frame(body)
        controls.pack(fill="x", pady=(12, 0))
        self.inspect_btn = ttk.Button(controls, text="فحص البيانات", command=self.inspect, style="Primary.TButton", state="disabled")
        self.inspect_btn.pack(side="right")
        self.update_btn = ttk.Button(controls, text="تحديث وإنشاء نسخة جديدة", command=self.perform_update, style="Success.TButton", state="disabled")
        self.update_btn.pack(side="right", padx=(0, 8))
        self.status = ttk.Label(controls, text="", foreground=MUTED, anchor="e")
        self.status.pack(side="right", fill="x", expand=True, padx=14)

        report_card = self.card(body)
        report_card.pack(fill="both", expand=True, pady=(12, 0))
        ttk.Label(report_card, text="3. التقرير", style="CardTitle.TLabel", anchor="e").pack(fill="x")
        self.report_note = ttk.Label(report_card, text="سيظهر تقرير الفحص والتحديث هنا.", style="Card.TLabel", anchor="e")
        self.report_note.pack(fill="x", pady=(4, 7))
        self.tabs = ttk.Notebook(report_card)
        self.tabs.pack(fill="both", expand=True)
        self.summary = self.text_tab("الملخص")
        self.base_missing = self.text_tab("غير مرتبطة في الأساسي")
        self.gis_missing = self.text_tab("غير مرتبطة في GIS")
        self.dupes = self.text_tab("الأكواد المكررة")
        self.change_tree = self.change_tab()

    def text_tab(self, title):
        tab = ttk.Frame(self.tabs, style="Card.TFrame", padding=8)
        self.tabs.add(tab, text=title)
        toolbar = ttk.Frame(tab, style="Card.TFrame")
        toolbar.pack(fill="x", pady=(0, 4))
        widget = tk.Text(tab, height=8, wrap="word", relief="flat", bg="#FBFCFE", fg=INK, font=("Consolas", 10))
        ttk.Button(toolbar, text="نسخ المحتوى", command=lambda: self.copy(widget)).pack(side="right")
        widget.pack(fill="both", expand=True)
        widget.configure(state="disabled")
        return widget

    def change_tab(self):
        tab = ttk.Frame(self.tabs, style="Card.TFrame", padding=8)
        self.tabs.add(tab, text="تفاصيل التغييرات")
        columns = ("code", "field", "old", "new", "reason")
        tree = ttk.Treeview(tab, columns=columns, show="headings")
        labels = ["كود القطعة", "الحقل", "القيمة القديمة", "القيمة الجديدة", "السبب"]
        widths = [130, 210, 180, 180, 150]
        for column, label, width in zip(columns, labels, widths):
            tree.heading(column, text=label, anchor="e")
            tree.column(column, width=width, anchor="e")
        bar = ttk.Scrollbar(tab, orient="vertical", command=tree.yview)
        tree.configure(yscrollcommand=bar.set)
        tree.pack(side="left", fill="both", expand=True)
        bar.pack(side="right", fill="y")
        return tree

    def choose_path(self):
        return filedialog.askopenfilename(
            title="اختيار ملف",
            filetypes=[("ملفات Excel", "*.xlsx *.xls *.xlsm *.xlsb"), ("CSV", "*.csv *.tsv"), ("كل الملفات", "*.*")]
        )

    def pick_base(self):
        path = self.choose_path()
        if path:
            self.load("base", path)

    def pick_gis(self):
        path = self.choose_path()
        if path:
            self.load("gis", path)

    def load(self, target, path):
        try:
            self.status.configure(text="جارِ قراءة الملف...")
            self.update_idletasks()
            loaded = read_file(path)
        except DataError as error:
            messagebox.showerror("تعذر قراءة الملف", str(error))
            self.status.configure(text="")
            return
        if target == "base":
            self.base = loaded
            self.base_status.configure(text=f"{Path(path).name} — {len(loaded.rows):,} صف")
        else:
            self.gis = loaded
            self.gis_status.configure(text=f"{Path(path).name} — {len(loaded.rows):,} صف")
        self.status.configure(text="تم تحميل الملف بنجاح")
        if self.base and self.gis:
            self.build_mapping()

    def build_mapping(self):
        for item in self.map_area.inner.winfo_children():
            item.destroy()
        self.mapping_vars.clear()
        suggested, base_key, gis_key = defaults(self.base, self.gis)
        self.base_key.set(base_key or "")
        self.gis_key.set(gis_key or "")
        area = self.map_area.inner
        area.columnconfigure(0, weight=2)
        area.columnconfigure(1, weight=2)
        area.columnconfigure(2, weight=1)
        for column, title in enumerate(["الحقل في الملف الأساسي", "العمود المستخدم من GIS", "الحالة"]):
            ttk.Label(area, text=title, style="CardTitle.TLabel", anchor="e").grid(row=0, column=column, sticky="ew", padx=7, pady=(0, 5))

        key_row = ttk.Frame(area, style="Card.TFrame")
        key_row.grid(row=1, column=0, columnspan=3, sticky="ew", padx=7, pady=(0, 8))
        key_row.columnconfigure(1, weight=1)
        key_row.columnconfigure(3, weight=1)
        ttk.Label(key_row, text="كود القطعة في الأساسي", style="Card.TLabel").grid(row=0, column=0, padx=(0, 6))
        ttk.Combobox(key_row, textvariable=self.base_key, values=self.base.headers, state="readonly").grid(row=0, column=1, sticky="ew", padx=(0, 15))
        ttk.Label(key_row, text="كود القطعة في GIS", style="Card.TLabel").grid(row=0, column=2, padx=(0, 6))
        ttk.Combobox(key_row, textvariable=self.gis_key, values=self.gis.headers, state="readonly").grid(row=0, column=3, sticky="ew")

        options = [SKIP, *self.gis.headers]
        for row_no, (target, source) in enumerate(suggested.items(), start=2):
            value = tk.StringVar(value=source or SKIP)
            self.mapping_vars[target] = value
            ttk.Label(area, text=target, style="Card.TLabel", anchor="e").grid(row=row_no, column=0, sticky="ew", padx=7, pady=4)
            ttk.Combobox(area, textvariable=value, values=options, state="readonly").grid(row=row_no, column=1, sticky="ew", padx=7, pady=4)
            state = ttk.Label(area, style="Card.TLabel", anchor="e")
            state.grid(row=row_no, column=2, sticky="ew", padx=7, pady=4)

            def refresh(*_args, variable=value, label=state, base_header=target):
                if base_header not in self.base.headers:
                    label.configure(text="غير موجود في الأساسي", foreground=BAD)
                elif variable.get() == SKIP:
                    label.configure(text="لن يتم تحديثه", foreground=WARN)
                else:
                    label.configure(text="تم العثور عليه", foreground=OK)
            value.trace_add("write", refresh)
            refresh()

        self.map_note.configure(text="تم تركيب الأعمدة تلقائيًا. عدّل أي عمود من القائمة إذا كان ملف GIS مختلفًا.")
        self.inspect_btn.configure(state="normal")
        self.update_btn.configure(state="normal")
        self.clear_report()

    def selected_mapping(self):
        return {target: None if value.get() == SKIP else value.get() for target, value in self.mapping_vars.items()}

    def selection_valid(self):
        if not self.base_key.get() or not self.gis_key.get():
            messagebox.showerror("كود القطعة", "اختر عمود كود القطعة في الملفين.")
            return False
        return True

    def inspect(self):
        if not self.selection_valid():
            return
        try:
            base_copy = Table(self.base.path, self.base.headers, deepcopy(self.base.rows))
            gis_copy = Table(self.gis.path, self.gis.headers, deepcopy(self.gis.rows))
            report = apply_update(base_copy, gis_copy, self.base_key.get(), self.gis_key.get(), self.selected_mapping())
            self.show_report(report, "تقرير فحص قبل التنفيذ")
            self.status.configure(text="اكتمل الفحص. راجع التقرير قبل إنشاء النسخة الجديدة.")
        except DataError as error:
            messagebox.showerror("تعذر الفحص", str(error))

    def perform_update(self):
        if not self.selection_valid():
            return
        destination = output_path(self.base)
        if not messagebox.askyesno("تأكيد التحديث", f"سيتم إنشاء نسخة جديدة باسم:\n{destination.name}\n\nالملف الأصلي لن يتغير."):
            return
        try:
            self.status.configure(text="جارِ تحديث البيانات وحفظ النسخة الجديدة...")
            self.update_idletasks()
            report = apply_update(self.base, self.gis, self.base_key.get(), self.gis_key.get(), self.selected_mapping())
            saved = save(self.base, destination)
            self.show_report(report, "تقرير التحديث والحفظ")
            self.status.configure(text=f"تم الحفظ: {saved.name}")
            messagebox.showinfo("تم التحديث", f"تم إنشاء النسخة الجديدة بنجاح:\n{saved}")
        except (DataError, OSError) as error:
            messagebox.showerror("تعذر إنشاء النسخة", str(error))

    def clear_report(self):
        for widget in (self.summary, self.base_missing, self.gis_missing, self.dupes):
            widget.configure(state="normal")
            widget.delete("1.0", "end")
            widget.configure(state="disabled")
        for item in self.change_tree.get_children():
            self.change_tree.delete(item)
        self.report_note.configure(text="سيظهر تقرير الفحص والتحديث هنا.")

    def set_text(self, widget, content):
        widget.configure(state="normal")
        widget.delete("1.0", "end")
        widget.insert("1.0", content if content else "لا توجد بيانات.")
        widget.configure(state="disabled")

    def show_report(self, report, title):
        self.report_note.configure(text=title)
        self.set_text(self.summary, report_text(report))
        self.set_text(self.base_missing, "\n".join(report.base_unmatched))
        self.set_text(self.gis_missing, "\n".join(report.gis_unmatched))
        self.set_text(self.dupes, "\n".join([
            "مكررة في الملف الأساسي:", *report.base_duplicates,
            "", "مكررة في ملف GIS:", *report.gis_duplicates,
        ]))
        for item in self.change_tree.get_children():
            self.change_tree.delete(item)
        for parcel, field, old, new, reason in report.details:
            self.change_tree.insert("", "end", values=(parcel, field, self.show(old), self.show(new), reason))
        self.tabs.select(0)

    @staticmethod
    def show(value):
        return "فارغ" if blank(value) else str(value)

    def copy(self, widget):
        self.clipboard_clear()
        self.clipboard_append(widget.get("1.0", "end-1c"))
        self.update_idletasks()
        self.status.configure(text="تم نسخ المحتوى إلى الحافظة")


if __name__ == "__main__":
    App().mainloop()
