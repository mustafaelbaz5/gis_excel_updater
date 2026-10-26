from __future__ import annotations

from collections.abc import Iterable
from pathlib import Path
from typing import Any

from .exceptions import SpreadsheetValidationError
from .models import WorkbookTable
from .normalization import display_text, is_blank_source


class WorkbookStore:
    SUPPORTED_SUFFIXES = {".xlsx", ".xlsm", ".xls", ".xlsb", ".csv", ".tsv"}

    def read(self, path_value: str | Path) -> WorkbookTable:
        path = Path(path_value)
        if not path.exists():
            raise SpreadsheetValidationError("الملف المحدد غير موجود.")
        if path.suffix.lower() not in self.SUPPORTED_SUFFIXES:
            raise SpreadsheetValidationError("الصيغة غير مدعومة. اختر xlsx أو xls أو xlsm أو xlsb أو csv أو tsv.")
        return self._read_openpyxl(path) if path.suffix.lower() in {".xlsx", ".xlsm"} else self._read_tabular(path)

    def write_cell(self, table: WorkbookTable, row: dict[str, Any], header: str, value: Any) -> None:
        row[header] = value
        if table.is_native_excel:
            column = table.headers.index(header) + 1
            table.native_sheet.cell(row=row["__excel_row__"], column=column).value = value

    def save_as_new_file(self, table: WorkbookTable) -> Path:
        destination = self._available_output_path(table.path)
        if table.is_native_excel:
            table.native_workbook.save(destination)
            return destination

        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "البيانات المحدثة"
        sheet.append(table.headers)
        for row in table.rows:
            sheet.append([row.get(header) for header in table.headers])
        sheet.freeze_panes = "A2"
        workbook.save(destination)
        return destination

    @staticmethod
    def _available_output_path(source_path: Path) -> Path:
        suffix = source_path.suffix.lower() if source_path.suffix.lower() in {".xlsx", ".xlsm"} else ".xlsx"
        candidate = source_path.with_name(f"{source_path.stem}_محدث{suffix}")
        number = 2
        while candidate.exists():
            candidate = source_path.with_name(f"{source_path.stem}_محدث_{number}{suffix}")
            number += 1
        return candidate

    def _read_openpyxl(self, path: Path) -> WorkbookTable:
        from openpyxl import load_workbook

        workbook = load_workbook(path, keep_vba=path.suffix.lower() == ".xlsm", data_only=False)
        sheet = self._single_data_sheet(workbook.worksheets)
        matrix = [list(row) for row in sheet.iter_rows(values_only=True)]
        header_index = self._detect_header_index(matrix)
        headers = self._unique_headers(matrix[header_index])
        return WorkbookTable(
            path=path,
            headers=headers,
            rows=self._records(matrix, headers, header_index),
            native_workbook=workbook,
            native_sheet=sheet,
            header_row=header_index + 1,
        )

    def _read_tabular(self, path: Path) -> WorkbookTable:
        import pandas as pd

        try:
            if path.suffix.lower() in {".csv", ".tsv"}:
                frame = pd.read_csv(path, header=None, sep="\t" if path.suffix.lower() == ".tsv" else None, engine="python")
            else:
                workbook = pd.ExcelFile(path)
                if len(workbook.sheet_names) != 1:
                    raise SpreadsheetValidationError("الملف يحتوي على أكثر من شيت. استخدم ملفًا به شيت بيانات واحد.")
                frame = pd.read_excel(path, header=None)
        except SpreadsheetValidationError:
            raise
        except Exception as error:
            raise SpreadsheetValidationError(f"تعذر قراءة الملف: {error}") from error

        matrix = frame.where(frame.notna(), None).values.tolist()
        header_index = self._detect_header_index(matrix)
        return WorkbookTable(path, self._unique_headers(matrix[header_index]), self._records(matrix, self._unique_headers(matrix[header_index]), header_index))

    @staticmethod
    def _single_data_sheet(sheets: Iterable[Any]) -> Any:
        candidates = []
        for sheet in sheets:
            if sheet.sheet_state != "visible":
                continue
            sample = [list(row) for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 20), values_only=True)]
            if any(sum(not is_blank_source(value) for value in row) >= 2 for row in sample):
                candidates.append(sheet)
        if not candidates:
            raise SpreadsheetValidationError("لم يتم العثور على شيت بيانات صالح.")
        if len(candidates) != 1:
            raise SpreadsheetValidationError("الملف يحتوي على أكثر من شيت به بيانات. الأداة تعمل على شيت بيانات واحد فقط.")
        return candidates[0]

    @staticmethod
    def _detect_header_index(matrix: list[list[Any]]) -> int:
        if not matrix:
            raise SpreadsheetValidationError("الملف لا يحتوي على بيانات.")
        counts = [sum(not is_blank_source(value) for value in row) for row in matrix[:20]]
        best = max(range(len(counts)), key=counts.__getitem__)
        if counts[best] == 0:
            raise SpreadsheetValidationError("لم يتم العثور على صف عناوين للأعمدة.")
        return best

    @staticmethod
    def _unique_headers(raw_headers: list[Any]) -> list[str]:
        headers, seen = [], set()
        for index, value in enumerate(raw_headers, 1):
            original = display_text(value) or f"عمود_{index}"
            candidate, number = original, 2
            from .normalization import normalized_text
            while normalized_text(candidate) in seen:
                candidate = f"{original} ({number})"
                number += 1
            headers.append(candidate)
            seen.add(normalized_text(candidate))
        return headers

    @staticmethod
    def _records(matrix: list[list[Any]], headers: list[str], header_index: int) -> list[dict[str, Any]]:
        records = []
        for excel_row, values in enumerate(matrix[header_index + 1:], start=header_index + 2):
            if not any(not is_blank_source(value) for value in values):
                continue
            record = {header: values[index] if index < len(values) else None for index, header in enumerate(headers)}
            record["__excel_row__"] = excel_row
            records.append(record)
        return records
